"""Parse the raw university Workshop Schedule Excel workbook into records.

The workbook is a visual matrix, not a flat table. This module extracts the
structure programmatically (category columns, positions, schedule sections,
week numbers and the KEY legend) and turns it into intermediate records.
"""

import re
from dataclasses import dataclass, field

import openpyxl

DAY_NAMES = {
    "MONDAY": "MONDAY",
    "TUESDAY": "TUESDAY",
    "WEDNESDAY": "WEDNESDAY",
    "THURSDAY": "THURSDAY",
    "FRIDAY": "FRIDAY",
}

PERIOD_NAMES = {
    "MORNING": "MORNING",
    "AFTERNOON": "AFTERNOON",
}

GROUP_CODE_RE = re.compile(r"^[A-Z]{1,2}\d{0,2}$")
YEAR_RE = re.compile(r"\b(\d{4}/\d{4})\b")
SEMESTER_RE = re.compile(r"SEMESTER\s+([1-4])", re.IGNORECASE)
YEAR_OF_STUDY_RE = re.compile(r"\b(FIRST|SECOND|THIRD|FOURTH)\s+YEAR\b", re.IGNORECASE)

YEAR_OF_STUDY_MAP = {
    "FIRST": 1,
    "SECOND": 2,
    "THIRD": 3,
    "FOURTH": 4,
}


@dataclass
class WorkshopRecord:
    academic_year: str
    semester: int
    year_of_study: int | None
    schedule_section: str
    week_start: int
    week_end: int
    weeks: list = field(default_factory=list)
    group_code: str = ""
    workshop: str = ""
    position: int = 0
    day: str = ""
    time_period: str = ""
    venue: str = ""
    course_code: str = ""
    source_cells: list = field(default_factory=list)

    @property
    def week_label(self):
        if self.week_start == self.week_end:
            return str(self.week_start)
        return f"{self.week_start}-{self.week_end}"


@dataclass
class ParsedSection:
    name: str
    start_row: int
    end_row: int
    weeks: list = field(default_factory=list)
    week_by_row: dict = field(default_factory=dict)
    cell_bags: dict = field(default_factory=dict)


@dataclass
class ParsedWorkbook:
    sheet_name: str = "Sheet1"
    format: str = "matrix"
    academic_year: str = ""
    semester: int | None = None
    year_of_study: int | None = None
    key: dict = field(default_factory=dict)
    positions: dict = field(default_factory=dict)
    categories: dict = field(default_factory=dict)
    sections: list = field(default_factory=list)
    records: list = field(default_factory=list)
    unrecognized_cells: list = field(default_factory=list)
    invalid_groups: list = field(default_factory=list)
    unknown_keys: list = field(default_factory=list)
    duplicate_keys: list = field(default_factory=list)
    errors: list = field(default_factory=list)


def load_workbook(path):
    return openpyxl.load_workbook(path, data_only=True)


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_int(value) -> int | None:
    if value is None:
        return None
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return None


def _upper(value) -> str:
    return _cell_text(value).upper()


def detect_format(path):
    """Return 'matrix' when the sheet looks like the raw schedule workbook."""
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    if ws.max_row is None or ws.max_column is None:
        return "flat"
    markers = set()
    for row in ws.iter_rows():
        for cell in row:
            v = _upper(cell.value)
            if not v:
                continue
            if v in ("GROUPS", "POSITION", "KEY") or v.startswith("SCHEDULE"):
                markers.add(v)
    return "matrix" if {"GROUPS", "POSITION"} <= markers else "flat"


def detect_metadata(texts: list[str]) -> tuple[str, int | None, int | None]:
    """Return (academic_year, semester, year_of_study) from worksheet text."""
    academic_year = ""
    semester = None
    year_of_study = None
    for text in texts:
        m = YEAR_RE.search(text)
        if m and not academic_year:
            academic_year = m.group(1)
        m = SEMESTER_RE.search(text)
        if m and semester is None:
            semester = int(m.group(1))
        m = YEAR_OF_STUDY_RE.search(text)
        if m and year_of_study is None:
            year_of_study = YEAR_OF_STUDY_MAP.get(m.group(1).upper())
    return academic_year, semester, year_of_study


def parse_key_description(text: str) -> list[tuple[str, str]]:
    """Turn a KEY description into [(day, period), ...] tuples."""
    parts = re.split(r"\s*(?:and|&)\s*", text)
    pairs = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        tokens = part.split()
        if len(tokens) < 2:
            continue
        day = DAY_NAMES.get(tokens[0].upper())
        period = PERIOD_NAMES.get(tokens[-1].upper())
        if day and period:
            pairs.append((day, period))
    return pairs


def _merge_map(ws, row_idx):
    """Map a column to the label/category for a merged row region."""
    mapping = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row_idx <= rng.max_row:
            value = _cell_text(ws.cell(row=row_idx, column=rng.min_col).value)
            if value:
                for col in range(rng.min_col, rng.max_col + 1):
                    mapping[col] = value
    return mapping


def _find_marker_row(ws, marker, max_row):
    for row_idx in range(1, max_row + 1):
        v = _upper(ws.cell(row=row_idx, column=1).value)
        if v == marker:
            return row_idx
    return None


def _find_categories_row(ws, groups_row):
    best_row, best_count = None, 0
    for row_idx in range(1, groups_row):
        count = sum(
            1 for cell in ws[row_idx] if _cell_text(cell.value)
        )
        if count > best_count:
            best_row, best_count = row_idx, count
    return best_row


def _categorise(parsed: ParsedWorkbook, ws, max_row):
    groups_row = _find_marker_row(ws, "GROUPS", max_row)
    if groups_row is None:
        parsed.errors.append("GROUPS header row not found")
        return
    position_row = _find_marker_row(ws, "POSITION", max_row)
    if position_row is None:
        parsed.errors.append("POSITION header row not found")
        return

    cat_row = _find_categories_row(ws, groups_row)
    if cat_row is None:
        parsed.errors.append("Workshop category row not found")
        return

    parsed.categories = _merge_map(ws, cat_row)
    parsed.positions = {}
    for col in range(1, ws.max_column + 1):
        pos = _cell_int(ws.cell(row=position_row, column=col).value)
        if pos is not None:
            parsed.positions[col] = pos


def _extract_key(parsed: ParsedWorkbook, ws, max_row):
    key_row = _find_marker_row(ws, "KEY", max_row)
    if key_row is None:
        parsed.errors.append("KEY legend not found")
        return
    for row_idx in range(key_row, max_row + 1):
        row_pairs = []
        for col in range(1, ws.max_column + 1):
            num = _cell_int(ws.cell(row=row_idx, column=col).value)
            if num is None or not (1 <= num <= 6):
                continue
            text = ""
            for c2 in range(col + 1, min(col + 3, ws.max_column) + 1):
                text = _cell_text(ws.cell(row=row_idx, column=c2).value)
                if text:
                    break
            if text:
                row_pairs.append((num, parse_key_description(text)))
        for num, pairs in row_pairs:
            existing = parsed.key.setdefault(num, [])
            for pair in pairs:
                if pair not in existing:
                    existing.append(pair)


def _extract_sections(parsed: ParsedWorkbook, ws, max_row):
    key_row = _find_marker_row(ws, "KEY", max_row) or max_row + 1
    section_rows = []
    for row_idx in range(1, key_row):
        v = _upper(ws.cell(row=row_idx, column=1).value)
        if v.startswith("SCHEDULE"):
            section_rows.append(row_idx)
    if not section_rows:
        parsed.errors.append("No SCHEDULE sections found")
        return
    section_rows.append(key_row)
    for i, start_row in enumerate(section_rows[:-1]):
        end_row = section_rows[i + 1] - 1
        name = _cell_text(ws.cell(row=start_row, column=1).value) or f"Section {i + 1}"
        section = ParsedSection(name=name, start_row=start_row, end_row=end_row)
        for row_idx in range(start_row, end_row + 1):
            week = _cell_int(ws.cell(row=row_idx, column=2).value)
            if week is not None and week > 0:
                section.weeks.append(week)
                section.week_by_row[row_idx] = week
        section.weeks = sorted(section.weeks)
        if section.weeks:
            parsed.sections.append(section)


def _split_codes(value: str) -> list[str]:
    codes = []
    for part in value.split(","):
        part = part.strip()
        if part:
            codes.append(part)
    return codes


def _collect_cells(parsed: ParsedWorkbook, ws, section: ParsedSection):
    for row_idx in range(section.start_row, section.end_row + 1):
        week = section.week_by_row.get(row_idx)
        for col in range(1, ws.max_column + 1):
            if col <= 2:
                continue
            raw = _cell_text(ws.cell(row=row_idx, column=col).value)
            if not raw:
                continue
            coord = f"{openpyxl.utils.get_column_letter(col)}{row_idx}"
            if week is None:
                parsed.unrecognized_cells.append(f"{coord}='{raw}' (no week number)")
                continue
            cat = parsed.categories.get(col)
            pos = parsed.positions.get(col)
            if cat is None or pos is None:
                parsed.unrecognized_cells.append(
                    f"{coord}='{raw}' (outside workshop columns)"
                )
                continue
            codes = _split_codes(raw)
            if not codes:
                parsed.invalid_groups.append(f"{coord}='{raw}'")
                continue
            if pos not in parsed.key:
                parsed.unknown_keys.append(f"{coord}: position {pos} has no KEY")
                continue
            for code in codes:
                if not GROUP_CODE_RE.match(code):
                    parsed.invalid_groups.append(f"{coord}='{raw}' -> '{code}'")
                    continue
                key = (section.name, cat, pos, code)
                bag = section.cell_bags.setdefault(key, [[], []])
                if week not in bag[0]:
                    bag[0].append(week)
                if coord not in bag[1]:
                    bag[1].append(coord)


def _to_runs(weeks):
    runs = []
    start = prev = weeks[0]
    for w in weeks[1:]:
        if w == prev + 1:
            prev = w
        else:
            runs.append((start, prev))
            start = prev = w
    runs.append((start, prev))
    return runs


def _build_records(parsed: ParsedWorkbook, ws):
    for section in parsed.sections:
        _collect_cells(parsed, ws, section)
        for (section_name, cat, pos, code), (weeks, cells) in section.cell_bags.items():
            week_list = sorted(set(weeks))
            for week_start, week_end in _to_runs(week_list):
                for day, period in parsed.key[pos]:
                    parsed.records.append(
                        WorkshopRecord(
                            academic_year=parsed.academic_year,
                            semester=parsed.semester or 0,
                            year_of_study=parsed.year_of_study,
                            schedule_section=section_name,
                            week_start=week_start,
                            week_end=week_end,
                            weeks=list(range(week_start, week_end + 1)),
                            group_code=code,
                            workshop=cat,
                            position=pos,
                            day=day,
                            time_period=period,
                            venue="",
                            course_code=cat,
                            source_cells=cells,
                        )
                    )


def _find_duplicates(parsed: ParsedWorkbook):
    seen = {}
    for rec in parsed.records:
        key = (
            rec.academic_year,
            rec.semester,
            rec.course_code,
            rec.group_code,
            rec.day,
            rec.time_period,
            rec.position,
            rec.schedule_section,
            rec.week_start,
            rec.week_end,
        )
        seen.setdefault(key, []).append(rec)
    for key, recs in seen.items():
        if len(recs) > 1:
            parsed.duplicate_keys.append(
                f"{key[2]} {key[3]} {key[4]} {key[5]} {key[6]} "
                f"{key[7]} weeks {key[8]}-{key[9]} ({len(recs)} records)"
            )


def parse_workbook(path) -> ParsedWorkbook:
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    parsed = ParsedWorkbook(sheet_name=ws.title)
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_row == 0 or max_col == 0:
        parsed.errors.append("Workbook has no cells")
        return parsed

    texts = []
    for row in ws.iter_rows():
        for cell in row:
            text = _cell_text(cell.value)
            if text:
                texts.append(text)
    parsed.academic_year, parsed.semester, parsed.year_of_study = detect_metadata(
        texts
    )

    _categorise(parsed, ws, max_row)
    _extract_key(parsed, ws, max_row)
    _extract_sections(parsed, ws, max_row)
    _build_records(parsed, ws)
    _find_duplicates(parsed)
    return parsed